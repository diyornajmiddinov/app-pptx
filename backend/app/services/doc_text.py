"""
Plain text out of a document the user uploads on the website, for "presentation from my document":
.docx (Word XML), .pptx (python-pptx), .pdf (poppler's pdftotext), .txt / .md. Nothing is executed or rendered;
the text is cut to content.SOURCE_MAX characters for the prompt.
"""
import io
import re
import subprocess
import tempfile
import zipfile
from xml.etree import ElementTree

from pptx import Presentation

MAX_FILE = 15 * 1024 * 1024
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class DocError(Exception):
    """str(e): docType (not a supported file) or docEmpty (no text found, e.g. a scanned PDF)."""


def _docx(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        info = z.getinfo("word/document.xml")
        if info.file_size > 50 * 1024 * 1024:  # a zip bomb
            raise DocError("docType")
        root = ElementTree.fromstring(z.read(info))
    paras = ("".join(t.text or "" for t in p.iter(f"{W}t")) for p in root.iter(f"{W}p"))
    return "\n".join(p for p in paras if p.strip())


def _pptx(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
    return "\n".join(out)


def _pdf(data: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".pdf") as f:
        f.write(data); f.flush()
        r = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", "-l", "60", f.name, "-"],
                           capture_output=True, timeout=60, check=False)
    if r.returncode:
        raise DocError("docType")
    return r.stdout.decode("utf-8", errors="replace")


def extract(data: bytes, filename: str) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if ext == "docx" or (data[:2] == b"PK" and ext not in ("pptx",)):
            text = _docx(data)
        elif ext == "pptx":
            text = _pptx(data)
        elif ext == "pdf" or data[:5] == b"%PDF-":
            text = _pdf(data)
        elif ext in ("txt", "md", ""):
            text = data.decode("utf-8")
        else:
            raise DocError("docType")
    except DocError:
        raise
    except Exception as e:  # a broken file: zipfile, XML, python-pptx and pdftotext each raise their own types
        raise DocError("docType") from e
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    if len(text) < 50:
        raise DocError("docEmpty")
    return text
