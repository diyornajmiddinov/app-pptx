"""
Presentation editor (desktop only): edits ANY .pptx in place with python-pptx, so everything the editor
does not touch (animations, layouts, themes, charts, fonts) stays exactly as it was in the file.

data/<user_id>/edits/<doc_id>/
    meta.json            {"id", "name", "source", "created", "updated", "versions": [1, 2, ...], "cur": 2}
    v<N>.pptx            every saved state (undo / redo walk this list; the oldest are dropped after MAX_VERSIONS)
    layers/<key>.png     background layer of a slide (everything the browser does not draw), by content key
    media/<sha1>.<ext>   pictures of the file, served to the browser
    uploads/<id>.<ext>   images the user added

describe() turns a file into the model the browser draws (pptx_model.py); apply() runs a list of operations
on it. Text, moves and styles are drawn by the browser at once and saved in the background, so the user never waits.
Positions are in EMU (914400 per inch), as in the file.
"""
import copy
import io
import json
import re
import secrets
import shutil
import subprocess
import tempfile
import time
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from app.core.config import DATA_DIR
from app.core.locks import machine_lock
from app.services import pptx_model

DATA = DATA_DIR
MAX_FILE = 40 * 1024 * 1024          # uploaded .pptx
MAX_UNZIPPED = 300 * 1024 * 1024     # guard against zip bombs
MAX_IMAGE = 10 * 1024 * 1024
MAX_VERSIONS = 30
IMPORTABLE = (".pptx", ".ppt", ".odp", ".pps", ".ppsx")

ALIGN = {PP_ALIGN.LEFT: "left", PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right", PP_ALIGN.JUSTIFY: "justify"}
ALIGN_SET = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}


class EditorError(Exception):
    """A request the editor refuses; the message key is translated by the editor page."""


# ---------------------------------------------------------------- storage

def _root(user_id: int) -> Path:
    return DATA / str(user_id) / "edits"


def _folder(user_id: int, doc_id: str) -> Path:
    if not re.fullmatch(r"[a-f0-9]{16}", str(doc_id)):
        raise EditorError("notFound")
    f = _root(user_id) / doc_id
    if not (f / "meta.json").exists():
        raise EditorError("notFound")
    return f


def _meta(folder: Path) -> dict:
    return json.loads((folder / "meta.json").read_text(encoding="utf-8"))


def _write_meta(folder: Path, meta: dict):
    tmp = folder / "meta.json.tmp"
    tmp.write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    tmp.replace(folder / "meta.json")


def _file(folder: Path, version: int) -> Path:
    return folder / f"v{version}.pptx"


def clean_name(name: str) -> str:
    name = Path(name or "taqdimot").name
    stem = re.sub(r"[\x00-\x1f/\\:*?\"<>|]+", "_", Path(name).stem).strip() or "taqdimot"
    return stem[:80]


def _check_zip(data: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
                raise EditorError("tooBig")
    except zipfile.BadZipFile as e:
        raise EditorError("badFile") from e


def _soffice(args: list[str], cwd: Path, timeout: int = 120):
    with machine_lock("office"), tempfile.TemporaryDirectory(prefix="lo-") as profile:  # one LibreOffice on the server at a time
        subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore", *args],
                       cwd=cwd, check=True, timeout=timeout, capture_output=True)


def import_file(user_id: int, data: bytes, filename: str, source: str = "upload") -> str:
    """Store an uploaded presentation as a new editor document. .ppt / .odp are converted to .pptx first."""
    if len(data) > MAX_FILE:
        raise EditorError("tooBig")
    ext = Path(filename or "").suffix.lower() or ".pptx"
    if ext not in IMPORTABLE:
        raise EditorError("badFile")
    if ext != ".pptx":  # legacy / OpenDocument formats: let LibreOffice convert them
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / f"in{ext}"; src.write_bytes(data)
            try:
                _soffice(["--convert-to", "pptx", "--outdir", tmp, str(src)], Path(tmp))
                data = (Path(tmp) / "in.pptx").read_bytes()
            except (subprocess.SubprocessError, FileNotFoundError) as e:
                raise EditorError("badFile") from e
    _check_zip(data)
    try:
        Presentation(io.BytesIO(data))
    except Exception as e:
        raise EditorError("badFile") from e
    doc_id = secrets.token_hex(8)
    folder = _root(user_id) / doc_id
    folder.mkdir(parents=True)
    _file(folder, 1).write_bytes(data)
    now = int(time.time())
    _write_meta(folder, {"id": doc_id, "name": clean_name(filename), "source": source, "created": now, "updated": now,
                         "versions": [1], "cur": 1})
    return doc_id


def list_docs(user_id: int) -> list[dict]:
    root, out = _root(user_id), []
    if root.exists():
        for m in root.glob("*/meta.json"):
            d = json.loads(m.read_text(encoding="utf-8"))
            out.append({k: d[k] for k in ("id", "name", "source", "created", "updated")})
    return sorted(out, key=lambda d: d["updated"], reverse=True)


def delete_doc(user_id: int, doc_id: str):
    shutil.rmtree(_folder(user_id, doc_id))


def current_file(user_id: int, doc_id: str) -> tuple[Path, str]:
    folder = _folder(user_id, doc_id); meta = _meta(folder)
    return _file(folder, meta["cur"]), meta["name"]


def rename(user_id: int, doc_id: str, name: str) -> dict:
    folder = _folder(user_id, doc_id); meta = _meta(folder)
    meta["name"] = clean_name(re.sub(r"[/\\]", "_", name) + ".pptx"); _write_meta(folder, meta)  # a name, not a path: keep what follows a slash
    return describe(user_id, doc_id)


def save_upload(user_id: int, doc_id: str, data: bytes) -> str:
    """An image for add_image / image ops. Checked with Pillow; returns its id."""
    folder = _folder(user_id, doc_id)
    if len(data) > MAX_IMAGE:
        raise EditorError("tooBig")
    try:
        with Image.open(io.BytesIO(data)) as im:
            fmt = (im.format or "").lower()
            im.verify()
    except Exception as e:
        raise EditorError("badImage") from e
    ext = {"jpeg": "jpg", "png": "png", "gif": "gif", "bmp": "bmp"}.get(fmt)
    if not ext:  # e.g. webp: convert, PowerPoint does not read it
        with Image.open(io.BytesIO(data)) as im:
            buf = io.BytesIO(); im.convert("RGBA").save(buf, "PNG"); data, ext = buf.getvalue(), "png"
    up = folder / "uploads"; up.mkdir(exist_ok=True)
    uid = secrets.token_hex(8)
    (up / f"{uid}.{ext}").write_bytes(data)
    return uid


def _upload_path(folder: Path, upload_id: str | None) -> Path:
    if not re.fullmatch(r"[a-f0-9]{16}", str(upload_id or "")):
        raise EditorError("badImage")
    found = list((folder / "uploads").glob(f"{upload_id}.*")) if (folder / "uploads").exists() else []
    if not found:
        raise EditorError("badImage")
    return found[0]


# ---------------------------------------------------------------- reading + background layer

def describe(user_id: int, doc_id: str) -> dict:
    """The editor model: per slide, the drawable elements (pptx_model) and the key of its background picture."""
    folder = _folder(user_id, doc_id); meta = _meta(folder)
    prs = Presentation(str(_file(folder, meta["cur"])))
    media = folder / "media"; media.mkdir(exist_ok=True)
    slides = []
    for i, s in enumerate(prs.slides):
        elements, ids, bg = pptx_model.slide_model(prs, s)
        for e in elements:  # pictures are served from media/<sha1>.<ext>
            img = e.pop("img", None)
            if img:
                f = media / f"{img['sha1']}.{img['ext']}"
                if not f.exists():
                    f.write_bytes(img["blob"])
                e["img"] = {"src": f.name, "crop": img["crop"]}
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide and s.notes_slide.notes_text_frame else ""
        slides.append({"n": i + 1, "hidden": s._element.get("show") == "0", "bg": bg, "notes": notes,
                       "layer": pptx_model.background_key(s, ids), "elements": elements})
    vs, cur = meta["versions"], meta["cur"]
    return {"id": doc_id, "name": meta["name"], "version": cur, "can_undo": vs.index(cur) > 0, "can_redo": vs.index(cur) < len(vs) - 1,
            "width": prs.slide_width, "height": prs.slide_height, "slides": slides}


def media_file(user_id: int, doc_id: str, name: str) -> Path:
    if not re.fullmatch(r"[a-f0-9]{40}\.(png|jpg|gif|bmp|webp|svg)", name or ""):
        raise EditorError("notFound")
    f = _folder(user_id, doc_id) / "media" / name
    if not f.exists():
        raise EditorError("notFound")
    return f


def layer_file(user_id: int, doc_id: str, key: str) -> Path:
    """Background picture of a slide. Rendered on first request for every slide of the current version
    whose key is missing; editing elements never changes a key, so this happens rarely."""
    if not re.fullmatch(r"[a-f0-9]{20}", key or ""):
        raise EditorError("notFound")
    folder = _folder(user_id, doc_id)
    out = folder / "layers"; out.mkdir(exist_ok=True)
    f = out / f"{key}.png"
    if f.exists():
        return f
    with machine_lock("layers"):
        if not f.exists():
            _render_layers(folder, out)
    if not f.exists():
        raise EditorError("notFound")  # an older version's key: the page shows the plain background colour
    return f


def _render_layers(folder: Path, out: Path):
    meta = _meta(folder)
    prs = Presentation(str(_file(folder, meta["cur"])))
    keys = []
    for s in prs.slides:
        _, ids, _ = pptx_model.slide_model(prs, s)
        keys.append(pptx_model.background_key(s, ids))
        pptx_model.strip(s._element, ids)
        if "show" in s._element.attrib:  # hidden slides are skipped by the PDF export
            del s._element.attrib["show"]
    if all((out / f"{k}.png").exists() for k in keys):
        return
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "layers.pptx"; prs.save(str(src))
        _soffice(["--convert-to", "pdf", "--outdir", tmp, str(src)], Path(tmp))
        subprocess.run(["pdftoppm", "-png", "-scale-to-x", "1600", "-scale-to-y", "-1", str(Path(tmp) / "layers.pdf"), str(Path(tmp) / "s")],
                       check=True, timeout=180, capture_output=True)
        pages = sorted(Path(tmp).glob("s-*.png"), key=lambda f: int(f.stem.split("-")[1]))
        for k, page in zip(keys, pages, strict=False):
            shutil.move(page, out / f"{k}.png")


# ---------------------------------------------------------------- operations

def _slide(prs, n):
    try:
        n = int(n)
    except (TypeError, ValueError) as e:
        raise EditorError("badOp") from e
    if not 1 <= n <= len(prs.slides):
        raise EditorError("badOp")
    return prs.slides[n - 1]


def _shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise EditorError("badOp")


def _emu(v) -> Emu:
    return Emu(round(float(v)))


def _set_para_text(p, text: str):
    if p.text == text:
        return
    runs = p.runs
    for el in p._p.findall(qn("a:br")) + p._p.findall(qn("a:fld")):  # breaks / fields would stay behind the new text
        p._p.remove(el)
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            p._p.remove(r._r)
    else:
        p.add_run().text = text


def _set_text(tf, texts: list[str]):
    """Replace the text paragraph by paragraph, keeping each paragraph's formatting (first run wins)."""
    texts = [str(t) for t in texts] or [""]
    paras = list(tf.paragraphs)
    for i, text in enumerate(texts):
        if i < len(paras):
            p = paras[i]
        else:  # new paragraph: a copy of the last one, so it looks the same
            new = copy.deepcopy(paras[-1]._p)
            paras[-1]._p.addnext(new); paras.append(None)
            p = tf.paragraphs[i]
        _set_para_text(p, text)
    for p in list(tf.paragraphs)[len(texts):]:
        p._p.getparent().remove(p._p)


def _style(tf, o: dict):
    for p in tf.paragraphs:
        if o.get("align") in ALIGN_SET:
            p.alignment = ALIGN_SET[o["align"]]
        for f in [r.font for r in p.runs] + [p.font]:
            if o.get("size"):
                f.size = Pt(max(4, min(400, float(o["size"]))))
            for k in ("bold", "italic", "underline"):
                if k in o:
                    setattr(f, k, bool(o[k]))
            if o.get("color"):
                f.color.rgb = RGBColor.from_string(str(o["color"]).lstrip("#")[:6])
            if o.get("font"):
                f.name = str(o["font"])[:60]


def _rpr(base, r: dict):
    """A run's rPr written from the browser's run (explicit values), keeping the original's language / link."""
    rpr = etree.SubElement(etree.Element("x"), qn("a:rPr"))
    if base is not None:
        for k in ("lang", "altLang", "dirty", "kern", "spc", "cap"):
            if base.get(k) is not None:
                rpr.set(k, base.get(k))
    if r.get("sz"):
        rpr.set("sz", str(round(max(1, min(4000, float(r["sz"]))) * 100)))
    rpr.set("b", "1" if r.get("b") else "0")
    rpr.set("i", "1" if r.get("i") else "0")
    rpr.set("u", "sng" if r.get("u") else "none")
    rpr.set("strike", "sngStrike" if r.get("s") else "noStrike")
    if r.get("baseline"):
        rpr.set("baseline", str(int(float(r["baseline"]) * 1000)))
    if r.get("color"):
        fill = etree.SubElement(rpr, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", re.sub(r"[^0-9A-Fa-f]", "", str(r["color"]))[:6].upper().ljust(6, "0"))
    if r.get("font"):
        for tag in ("a:latin", "a:ea", "a:cs"):
            etree.SubElement(rpr, qn(tag)).set("typeface", str(r["font"])[:60])
    if base is not None:
        for keep in ("a:hlinkClick", "a:hlinkMouseOver"):
            for el in base.findall(qn(keep)):
                rpr.append(copy.deepcopy(el))
    return rpr


def _write_paragraphs(tx_body, paras: list[dict]):
    """Replace all paragraphs of a txBody with the browser's paragraphs [{src, align, bullet?, runs:[{t|br, sz, b...}]}].
    Each new paragraph starts from the original paragraph it came from (src), so spacing, levels and bullets stay."""
    old = tx_body.findall(qn("a:p"))
    anchor = old[0] if old else None
    new_ps = []
    for pd in (paras or [{"runs": []}])[:500]:
        src = pd.get("src")
        base = old[src] if isinstance(src, int) and 0 <= src < len(old) else (old[-1] if old else None)
        p = etree.Element(qn("a:p"))
        bppr = base.find(qn("a:pPr")) if base is not None else None
        ppr = copy.deepcopy(bppr) if bppr is not None else None
        if pd.get("align") in pptx_model.ALIGN_BACK or "bullet" in pd:
            if ppr is None:
                ppr = etree.Element(qn("a:pPr"))
            if pd.get("align") in pptx_model.ALIGN_BACK:
                ppr.set("algn", pptx_model.ALIGN_BACK[pd["align"]])
            if "bullet" in pd:
                for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buBlip"):
                    for el in ppr.findall(qn(tag)):
                        ppr.remove(el)
                bu = etree.Element(qn("a:buChar") if pd["bullet"] else qn("a:buNone"))
                if pd["bullet"]:
                    bu.set("char", "•")
                    if ppr.get("marL") is None:
                        ppr.set("marL", "342900"); ppr.set("indent", "-342900")
                # schema order inside pPr: lnSpc, spcBef, spcAft, buClr.., buSz.., buFont, bu*, tabLst, defRPr, extLst
                after = [c for c in ppr if etree.QName(c).localname in ("tabLst", "defRPr", "extLst")]
                (after[0].addprevious(bu) if after else ppr.append(bu))
        if ppr is not None:
            p.append(ppr)
        base_rpr = next((r.find(qn("a:rPr")) for r in (base if base is not None else []) if r.tag == qn("a:r")), None)
        runs = pd.get("runs") or []
        for r in runs:
            if r.get("br"):
                br = etree.SubElement(p, qn("a:br")); br.append(_rpr(base_rpr, r))
            else:
                text = str(r.get("t", ""))
                if not text:
                    continue
                el = etree.SubElement(p, qn("a:r")); el.append(_rpr(base_rpr, r))
                etree.SubElement(el, qn("a:t")).text = text
        end = copy.deepcopy(_rpr(base_rpr, (runs[-1] if runs else pd.get("end")) or {}))
        end.tag = qn("a:endParaRPr")
        for k in ("b", "i", "u", "strike"):
            end.attrib.pop(k, None)
        p.append(end)
        new_ps.append(p)
    for p in new_ps:
        (anchor.addprevious(p) if anchor is not None else tx_body.append(p))
    for p in old:
        tx_body.remove(p)


def _shape_elements(tree):
    return [el for el in tree if el.tag in (qn("p:sp"), qn("p:pic"), qn("p:graphicFrame"), qn("p:grpSp"), qn("p:cxnSp"), qn("p:contentPart"))]


def _next_id(slide) -> int:
    return max([int(x) for x in slide._element.xpath(".//p:cNvPr/@id")] + [0]) + 1


def _cover_crop(pic, img_path: Path):
    """Crop the picture so the new image fills its box without stretching (like CSS object-fit: cover)."""
    with Image.open(img_path) as im:
        ia = im.width / im.height
    ba = (pic.width or 1) / (pic.height or 1)
    pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
    if ia > ba:
        pic.crop_left = pic.crop_right = (1 - ba / ia) / 2
    elif ia < ba:
        pic.crop_top = pic.crop_bottom = (1 - ia / ba) / 2


def _copy_rels(src_part, dst_part, elements):
    """Relationship ids used inside copied XML (images, links, media) are re-created in the target part."""
    mapping = {}
    for rel in src_part.rels.values():
        if rel.reltype.endswith(("/notesSlide", "/slideLayout")):
            continue
        mapping[rel.rId] = (dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external
                            else dst_part.relate_to(rel.target_part, rel.reltype))
    rns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for root in elements:
        for el in root.iter():
            for k, v in list(el.attrib.items()):
                if k.startswith(rns) and v in mapping:
                    el.set(k, mapping[v])


def _move_slide(prs, old: int, new: int):
    lst = prs.slides._sldIdLst
    items = list(lst)
    el = items[old]
    lst.remove(el)
    lst.insert(max(0, min(new, len(items) - 1)), el)


def _op(prs, folder: Path, o: dict):
    op = o.get("op")
    if op == "slide_add":
        layout = min(prs.slide_layouts, key=lambda l: len(l.placeholders))  # the emptiest layout ("Blank")
        prs.slides.add_slide(layout)
        _move_slide(prs, len(prs.slides) - 1, int(o.get("after", len(prs.slides) - 1)))
        return
    slide = _slide(prs, o.get("slide"))
    idx = int(o["slide"]) - 1
    if op == "slide_delete":
        if len(prs.slides) == 1:
            raise EditorError("lastSlide")
        el = list(prs.slides._sldIdLst)[idx]
        prs.part.drop_rel(el.get(qn("r:id"))); prs.slides._sldIdLst.remove(el)
    elif op == "slide_move":
        _move_slide(prs, idx, int(o["to"]) - 1)
    elif op == "slide_duplicate":
        new = prs.slides.add_slide(slide.slide_layout)
        for sh in list(new.shapes):
            sh._element.getparent().remove(sh._element)
        copies = [copy.deepcopy(el) for el in _shape_elements(slide.shapes._spTree)]
        for el in copies:
            new.shapes._spTree.append(el)
        bg = slide._element.cSld.find(qn("p:bg"))
        if bg is not None:
            bgc = copy.deepcopy(bg); new._element.cSld.insert(0, bgc); copies.append(bgc)
        _copy_rels(slide.part, new.part, copies)
        _move_slide(prs, len(prs.slides) - 1, idx + 1)
    elif op == "slide_hide":
        if o.get("hidden"):
            slide._element.set("show", "0")
        elif "show" in slide._element.attrib:
            del slide._element.attrib["show"]
    elif op == "bg":
        fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = RGBColor.from_string(str(o["color"]).lstrip("#")[:6])
    elif op == "notes":
        slide.notes_slide.notes_text_frame.text = str(o.get("text", ""))[:20000]
    elif op == "add_text":
        w, h = prs.slide_width, prs.slide_height
        tb = slide.shapes.add_textbox(_emu(o.get("x", w * 0.25)), _emu(o.get("y", h * 0.4)), _emu(o.get("w", w * 0.5)), _emu(o.get("h", h * 0.15)))
        tb.text_frame.word_wrap = True
        tb.text_frame.text = str(o.get("text", "Matn"))
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(float(o.get("size", 24)))
    elif op == "add_image":
        path = _upload_path(folder, o.get("upload"))
        with Image.open(path) as im:
            ratio = im.height / im.width
        w = int(prs.slide_width * 0.4); h = int(w * ratio)
        if h > prs.slide_height * 0.8:
            h = int(prs.slide_height * 0.8); w = int(h / ratio)
        slide.shapes.add_picture(str(path), _emu(o.get("x", (prs.slide_width - w) / 2)), _emu(o.get("y", (prs.slide_height - h) / 2)), Emu(w), Emu(h))
    else:
        sh = _shape(slide, o.get("id"))
        el = sh._element
        if op == "move":
            old_w, old_h = sh.width or 1, sh.height or 1
            for k, attr in (("x", "left"), ("y", "top"), ("w", "width"), ("h", "height")):
                if k in o:
                    setattr(sh, attr, _emu(max(0, float(o[k])) if k in ("w", "h") else float(o[k])))
            if getattr(sh, "has_table", False) and sh.has_table:  # a table keeps its proportions when resized
                fx, fy = (sh.width or 1) / old_w, (sh.height or 1) / old_h
                for col in sh.table.columns:
                    col.width = Emu(int(col.width * fx))
                for row in sh.table.rows:
                    row.height = Emu(int(row.height * fy))
        elif op == "rotate":
            sh.rotation = float(o.get("rot", 0)) % 360
        elif op == "text_runs":
            tx = el.find(qn("p:txBody"))
            if tx is None and el.tag == qn("p:sp"):  # a shape without text gets a centred text body
                tx = etree.fromstring(f'<p:txBody xmlns:p="{pptx_model.P}" xmlns:a="{pptx_model.A}">'
                                      '<a:bodyPr anchor="ctr" rtlCol="0"/><a:lstStyle/><a:p/></p:txBody>')
                ext = el.find(qn("p:extLst"))
                (ext.addprevious(tx) if ext is not None else el.append(tx))
            if tx is None:
                raise EditorError("badOp")
            _write_paragraphs(tx, o.get("paragraphs") or [])
        elif op == "cell_runs":
            if not getattr(sh, "has_table", False) or not sh.has_table:
                raise EditorError("badOp")
            r, c = int(o["r"]), int(o["c"])
            if not (0 <= r < len(sh.table.rows) and 0 <= c < len(sh.table.columns)):
                raise EditorError("badOp")
            _write_paragraphs(sh.table.cell(r, c)._tc.get_or_add_txBody(), o.get("paragraphs") or [])
        elif op == "text":
            if not sh.has_text_frame:
                raise EditorError("badOp")
            _set_text(sh.text_frame, list(o.get("paragraphs") or [""])[:500])
        elif op == "style":
            if not sh.has_text_frame:
                raise EditorError("badOp")
            _style(sh.text_frame, o)
        elif op == "fill":
            if o.get("color"):
                sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(str(o["color"]).lstrip("#")[:6])
            else:
                sh.fill.background()
        elif op == "cell":
            if not getattr(sh, "has_table", False) or not sh.has_table:
                raise EditorError("badOp")
            r, c = int(o["r"]), int(o["c"])
            if not (0 <= r < len(sh.table.rows) and 0 <= c < len(sh.table.columns)):
                raise EditorError("badOp")
            _set_text(sh.table.cell(r, c).text_frame, str(o.get("text", "")).split("\n"))
        elif op == "image":
            path = _upload_path(folder, o.get("upload"))
            blips = el.xpath(".//a:blip")
            if not blips:
                raise EditorError("badOp")
            _, rid = sh.part.get_or_add_image_part(str(path))
            blips[0].set(qn("r:embed"), rid)
            if hasattr(sh, "crop_left"):
                _cover_crop(sh, path)
        elif op == "delete":
            el.getparent().remove(el)
        elif op == "duplicate":
            new = copy.deepcopy(el)
            new.xpath("./*[1]/p:cNvPr")[0].set("id", str(_next_id(slide)))
            el.addnext(new)
            dup = next(s for s in slide.shapes if s._element is new)
            dup.left, dup.top = Emu((sh.left or 0) + 182880), Emu((sh.top or 0) + 182880)
        elif op == "order":
            tree = el.getparent(); shapes = _shape_elements(tree); i = shapes.index(el); where = o.get("where")
            if where == "front" and el is not shapes[-1]:
                shapes[-1].addnext(el)
            elif where == "back" and el is not shapes[0]:
                shapes[0].addprevious(el)
            elif where == "forward" and i < len(shapes) - 1:
                shapes[i + 1].addnext(el)
            elif where == "backward" and i > 0:
                shapes[i - 1].addprevious(el)
        elif op == "rename":
            sh.name = str(o.get("name", sh.name))[:80]
        else:
            raise EditorError("badOp")


def apply(user_id: int, doc_id: str, ops: list[dict], version: int, full: bool = True) -> dict:
    """Run the operations on the current version and save the result as a new version."""
    folder = _folder(user_id, doc_id); meta = _meta(folder)
    if version != meta["cur"]:
        raise EditorError("stale")  # changed in another window: the page reloads the document
    if not ops:
        return describe(user_id, doc_id)
    prs = Presentation(str(_file(folder, meta["cur"])))
    try:
        for o in ops[:200]:
            _op(prs, folder, o)
    except EditorError:
        raise
    except (KeyError, ValueError, TypeError, IndexError, AttributeError) as e:
        raise EditorError("badOp") from e
    nv = max(meta["versions"]) + 1
    prs.save(str(_file(folder, nv)))
    vs = [*meta["versions"][: meta["versions"].index(meta["cur"]) + 1], nv]  # a new change drops the redo branch
    for old in set(meta["versions"]) - set(vs[-MAX_VERSIONS:]):
        _file(folder, old).unlink(missing_ok=True)
    meta.update(versions=vs[-MAX_VERSIONS:], cur=nv, updated=int(time.time()))
    _write_meta(folder, meta)
    return describe(user_id, doc_id) if full else _state(meta)


def _state(meta: dict) -> dict:
    vs, cur = meta["versions"], meta["cur"]
    return {"version": cur, "can_undo": vs.index(cur) > 0, "can_redo": vs.index(cur) < len(vs) - 1}


def step(user_id: int, doc_id: str, delta: int, full: bool = True) -> dict:
    """Undo (-1) / redo (+1). The page has usually redrawn already (full=False: just the new version)."""
    folder = _folder(user_id, doc_id); meta = _meta(folder)
    i = meta["versions"].index(meta["cur"]) + delta
    if 0 <= i < len(meta["versions"]):
        meta.update(cur=meta["versions"][i], updated=int(time.time())); _write_meta(folder, meta)
    return describe(user_id, doc_id) if full else _state(meta)
