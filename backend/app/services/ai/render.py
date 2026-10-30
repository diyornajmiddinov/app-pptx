"""Draw a .pptx from content JSON with the Node engine (engine/fill.js)."""
import asyncio
import contextlib
import json
import os
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation

from app.core.config import ENGINE_DIR
from app.core.locks import machine_lock, machine_lock_async

TIMEOUT = 180  # s


def add_notes(path: Path, content: dict) -> Path:
    """Speaker notes (slide["notes"], a website option) go into the notes pages, one content slide per pptx slide.
    The engine draws no notes itself; nothing is touched when the deck has none."""
    notes = [s.get("notes") or "" for s in content.get("slides", [])]
    if not any(notes):
        return path
    prs = Presentation(str(path))
    if len(prs.slides) != len(notes):  # a template that splits or merges slides: better no notes than wrong ones
        return path
    for slide, text in zip(prs.slides, notes, strict=True):
        if text:
            slide.notes_slide.notes_text_frame.text = text
    prs.save(str(path))
    return path


def _content_file(content: dict) -> str:
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as f:
        json.dump(content, f, ensure_ascii=False)
    return f.name


async def render(content: dict, template_id: int, out_path: Path | str) -> Path:
    """Async: a cancelled job (time limit) kills node instead of leaving it running. One node on the server at a time
    (the "drawing" lock): the worker runs several jobs at once, but they mostly wait for the AI, drawing takes ~1 s."""
    src = _content_file(content)
    cmd = ["node", "fill.js", str(template_id), src, str(out_path)]
    try:
        async with machine_lock_async("drawing"):
            proc = await asyncio.create_subprocess_exec(*cmd, cwd=ENGINE_DIR, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
            try:
                out, err = await asyncio.wait_for(proc.communicate(), TIMEOUT)
            except BaseException:  # timeout or cancelled
                if proc.returncode is None:
                    proc.kill()
                    await proc.wait()
                raise
    finally:
        os.unlink(src)
    if proc.returncode:  # the same error type as subprocess.run(check=True): alerts.why explains it
        raise subprocess.CalledProcessError(proc.returncode, cmd, out.decode(errors="replace"), err.decode(errors="replace"))
    return await asyncio.to_thread(add_notes, Path(out_path), content)


def render_sync(content: dict, template_id: int, out_path: Path | str, lock: bool = True) -> Path:
    """For code that already runs in a thread (the editor's import). lock=False: the caller holds the "drawing" lock."""
    src = _content_file(content)
    try:
        with machine_lock("drawing") if lock else contextlib.nullcontext():
            subprocess.run(["node", "fill.js", str(template_id), src, str(out_path)], cwd=ENGINE_DIR, check=True, timeout=TIMEOUT,
                           capture_output=True, text=True)
    finally:
        os.unlink(src)
    return add_notes(Path(out_path), content)
