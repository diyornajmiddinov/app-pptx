"""
A referat / mustaqil ish as a Word file, laid out the way Uzbek universities ask for it:
A4, Times New Roman 14, line spacing 1.5, margins left 30 / right 15 / top 20 / bottom 20 mm, first-line indent
1.25 cm, justified text, page numbers in the footer centre (the title page counts but shows no number), every
structural part (KIRISH, each BOB, XULOSA, ADABIYOTLAR) on a new page, headings bold and centred without a full stop.
Font, spacing, margins and the page-number place can be changed per order (the teacher's own requirements).

    build(content, fmt, out_dir) -> (docx, pdf | None, pages)

Page numbers in the MUNDARIJA are real: the file is laid out once by LibreOffice (Liberation Serif has the metrics
of Times New Roman), the page of every heading is read from the PDF, and the file is written again with them.
The table of contents is a Word TOC field holding those numbers, so Word can also refresh it after the student edits.
"""
import logging
import re
import shutil
import subprocess
import tempfile
import threading
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

log = logging.getLogger("paper")
FONT = "Times New Roman"
BLANK = "_______________________"
DEFAULT_FORMAT = {"font_size": 14, "spacing": 1.5, "margin_left": 30, "margin_right": 15, "margin_top": 20, "margin_bottom": 20,
                  "page_numbers": "bottom_center", "toc_title": "mundarija"}
_office = threading.Lock()  # one LibreOffice at a time: 2 CPUs, 2 GB

ROMAN = ["I", "II", "III", "IV", "V", "VI"]
LABELS = {
    "uz": {"ministry": "O‘ZBEKISTON RESPUBLIKASI OLIY TA’LIM, FAN VA INNOVATSIYALAR VAZIRLIGI",
           "referat": "REFERAT", "mustaqil": "MUSTAQIL ISH", "subject": "Fan:", "topic": "Mavzu:",
           "done": "Bajardi:", "checked": "Qabul qildi:", "mundarija": "MUNDARIJA", "reja": "REJA",
           "intro": "KIRISH", "conclusion": "XULOSA", "references": "FOYDALANILGAN ADABIYOTLAR RO‘YXATI",
           "chapter": "{n} BOB.", "lang": "uz-Latn-UZ"},
    "uzc": {"ministry": "ЎЗБЕКИСТОН РЕСПУБЛИКАСИ ОЛИЙ ТАЪЛИМ, ФАН ВА ИННОВАЦИЯЛАР ВАЗИРЛИГИ",
            "referat": "РЕФЕРАТ", "mustaqil": "МУСТАҚИЛ ИШ", "subject": "Фан:", "topic": "Мавзу:",
            "done": "Бажарди:", "checked": "Қабул қилди:", "mundarija": "МУНДАРИЖА", "reja": "РЕЖА",
            "intro": "КИРИШ", "conclusion": "ХУЛОСА", "references": "ФОЙДАЛАНИЛГАН АДАБИЁТЛАР РЎЙХАТИ",
            "chapter": "{n} БОБ.", "lang": "uz-Cyrl-UZ"},
    "ru": {"ministry": "МИНИСТЕРСТВО ВЫСШЕГО ОБРАЗОВАНИЯ, НАУКИ И ИННОВАЦИЙ РЕСПУБЛИКИ УЗБЕКИСТАН",
           "referat": "РЕФЕРАТ", "mustaqil": "САМОСТОЯТЕЛЬНАЯ РАБОТА", "subject": "Дисциплина:", "topic": "Тема:",
           "done": "Выполнил(а):", "checked": "Проверил(а):", "mundarija": "СОДЕРЖАНИЕ", "reja": "ПЛАН",
           "intro": "ВВЕДЕНИЕ", "conclusion": "ЗАКЛЮЧЕНИЕ", "references": "СПИСОК ИСПОЛЬЗОВАННОЙ ЛИТЕРАТУРЫ",
           "chapter": "ГЛАВА {n}.", "lang": "ru-RU"},
    "en": {"ministry": "MINISTRY OF HIGHER EDUCATION, SCIENCE AND INNOVATIONS OF THE REPUBLIC OF UZBEKISTAN",
           "referat": "ESSAY", "mustaqil": "INDEPENDENT WORK", "subject": "Subject:", "topic": "Topic:",
           "done": "Done by:", "checked": "Checked by:", "mundarija": "CONTENTS", "reja": "PLAN",
           "intro": "INTRODUCTION", "conclusion": "CONCLUSION", "references": "REFERENCES",
           "chapter": "CHAPTER {n}.", "lang": "en-US"},
}
QUOTES = {"uz": "“{}”", "uzc": "«{}»", "ru": "«{}»", "en": "“{}”"}


def chapter_title(language: str, i: int, title: str) -> str:
    return f"{LABELS[language]['chapter'].format(n=ROMAN[i])} {title.upper()}"


def headings(content: dict) -> list[tuple[int, str]]:
    """(level, text) of every heading, in order: what the MUNDARIJA lists."""
    L = LABELS[content["meta"]["language"]]
    out = [(1, L["intro"])]
    for i, ch in enumerate(content["chapters"]):
        out.append((1, chapter_title(content["meta"]["language"], i, ch["title"])))
        out += [(2, f"{i + 1}.{j + 1}. {s['title']}") for j, s in enumerate(ch["sections"])]
    return [*out, (1, L["conclusion"]), (1, L["references"])]


# ---------------------------------------------------------------- low-level Word bits

def _fonts(rpr_owner, size: float | None = None, bold: bool | None = None, lang: str | None = None):
    """Times New Roman for every script (Latin, Cyrillic, complex), no theme font, black."""
    font = rpr_owner.font
    font.name, font.color.rgb = FONT, RGBColor(0, 0, 0)
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    rpr = rpr_owner.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rpr.insert(0, fonts)
    for a in list(fonts.attrib):
        if a.endswith("Theme"):
            del fonts.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), FONT)
    if lang:
        el = rpr.find(qn("w:lang"))
        if el is None:
            el = OxmlElement("w:lang"); rpr.append(el)
        el.set(qn("w:val"), lang)


def _field(paragraph, instr: str, shown: str = ""):
    """A Word field (PAGE, TOC ...) with `shown` as its last calculated result."""
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", shown), ("end", None)):
        run = paragraph.add_run()
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = f" {text} "; run._r.append(el)
        elif kind == "text":
            run.text = text
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); run._r.append(el)


def _fld_char(paragraph, kind: str, instr: str | None = None):
    run = paragraph.add_run()
    el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); run._r.append(el)
    if instr:
        run = paragraph.add_run()
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; run._r.append(it)


def _letter_spacing(run, points: float):
    """w:spacing goes before w:sz in the run properties (Word rejects the file otherwise)."""
    rpr = run._r.get_or_add_rPr()
    el = OxmlElement("w:spacing"); el.set(qn("w:val"), str(int(points * 20)))
    after = next((c for c in rpr if c.tag in (qn("w:w"), qn("w:kern"), qn("w:position"), qn("w:sz"), qn("w:szCs"))), None)
    if after is None:
        rpr.append(el)
    else:
        after.addprevious(el)


def _rich(paragraph, text: str):
    """Text with **bold** parts (the introduction's 'Mavzuning dolzarbligi.' and the like)."""
    for k, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(part).bold = bool(k % 2) or None


def _styles(doc, fmt: dict, lang: str):
    size, spacing = fmt["font_size"], fmt["spacing"]
    normal = doc.styles["Normal"]
    _fonts(normal, size, lang=lang)
    pf = normal.paragraph_format
    pf.alignment, pf.first_line_indent, pf.line_spacing = WD_ALIGN_PARAGRAPH.JUSTIFY, Cm(1.25), spacing
    pf.space_before = pf.space_after = Pt(0)
    pf.widow_control = True
    for name, level in (("Heading 1", 1), ("Heading 2", 2)):
        st = doc.styles[name]
        _fonts(st, size, bold=True)
        st.font.italic = False
        p = st.paragraph_format
        p.alignment, p.first_line_indent, p.line_spacing = WD_ALIGN_PARAGRAPH.CENTER, Cm(0), spacing
        p.space_before, p.space_after = Pt(0 if level == 1 else size), Pt(size)
        p.keep_with_next, p.page_break_before = True, level == 1
    width = Mm(210 - fmt["margin_left"] - fmt["margin_right"])
    for level in (1, 2):
        st = doc.styles.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = normal
        p = st.paragraph_format
        p.alignment, p.first_line_indent, p.left_indent = WD_ALIGN_PARAGRAPH.LEFT, Cm(0), Cm(0 if level == 1 else 0.75)
        p.space_after = Pt(4)
        p.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    return width


def _page_setup(section, fmt: dict):
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin, section.right_margin = Mm(fmt["margin_left"]), Mm(fmt["margin_right"])
    section.top_margin, section.bottom_margin = Mm(fmt["margin_top"]), Mm(fmt["margin_bottom"])
    section.header_distance = section.footer_distance = Mm(min(10, fmt["margin_bottom"] // 2, fmt["margin_top"] // 2))


def _plain(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, before: float = 0, indent_left: float = 0):
    pf = paragraph.paragraph_format
    pf.alignment, pf.first_line_indent, pf.line_spacing = align, Cm(0), 1.15
    pf.space_before, pf.left_indent = Pt(before), Cm(indent_left)
    return paragraph


# ---------------------------------------------------------------- the document

def _title_page(doc, meta: dict, fmt: dict):
    L, lang = LABELS[meta["language"]], meta["language"]
    size = fmt["font_size"]

    def line(text: str, bold=False, caps=False, before: float = 0, big: float | None = None, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0.0):
        p = _plain(doc.add_paragraph(), align, before, indent)
        run = p.add_run(text.upper() if caps else text)
        run.bold = bold or None
        if big:
            run.font.size = Pt(big)
        return p

    line(meta.get("ministry") or L["ministry"], bold=True, caps=True)
    line(meta.get("university") or BLANK, bold=True, caps=True, before=size)
    if meta.get("faculty"):
        line(meta["faculty"], before=size / 2)
    if meta.get("department"):
        line(meta["department"], before=size / 2 if not meta.get("faculty") else 0)
    doc_type = line(L[meta.get("doc_type") or "referat"], bold=True, before=size * 8, big=size + 10)
    _letter_spacing(doc_type.runs[0], 3)
    if meta.get("subject"):
        p = line("", before=size)
        p.add_run(f"{L['subject']} ").bold = True
        p.add_run(meta["subject"])
    p = line("", before=size / 2 if meta.get("subject") else size)
    p.add_run(f"{L['topic']} ").bold = True
    p.add_run(QUOTES[lang].format(meta["topic"])).bold = True
    right = 8.5  # cm from the left margin: the "done by / checked by" block sits on the right
    p = line("", align=WD_ALIGN_PARAGRAPH.LEFT, before=size * 9, indent=right)
    p.add_run(f"{L['done']} ").bold = True
    p.add_run(meta.get("group") or "")
    line(meta.get("student") or BLANK, align=WD_ALIGN_PARAGRAPH.LEFT, indent=right)
    p = line("", align=WD_ALIGN_PARAGRAPH.LEFT, before=size, indent=right)
    p.add_run(f"{L['checked']} ").bold = True
    line(meta.get("teacher") or BLANK, align=WD_ALIGN_PARAGRAPH.LEFT, indent=right)
    # the city and year sit at the bottom of the title page whatever the text above takes: its own footer
    foot = doc.sections[0].footer.paragraphs[0]
    _plain(foot)
    foot.add_run(" – ".join(x for x in (meta.get("city"), meta.get("year")) if x) or BLANK)


def _toc(doc, content: dict, fmt: dict, pages: dict[int, int] | None):
    L = LABELS[content["meta"]["language"]]
    # looks like a heading but is not one: a Heading style would put the MUNDARIJA into itself
    title = doc.add_paragraph()
    title.add_run(L[fmt.get("toc_title") or "mundarija"]).bold = True
    pf = title.paragraph_format
    pf.alignment, pf.first_line_indent, pf.space_after, pf.keep_with_next = WD_ALIGN_PARAGRAPH.CENTER, Cm(0), Pt(fmt["font_size"]), True
    items = headings(content)
    for k, (level, text) in enumerate(items):
        p = doc.add_paragraph(style=f"toc {level}")
        if k == 0:
            _fld_char(p, "begin", 'TOC \\o "1-2" \\h \\z \\u'); _fld_char(p, "separate")
        p.add_run(text if level == 2 else text.upper())
        p.add_run("\t" + str(pages.get(k, "") if pages else "00"))
        if k == len(items) - 1:
            _fld_char(p, "end")


def _heading(doc, text: str, level: int):
    return doc.add_paragraph(text, style=f"Heading {level}")


def _body(doc, content: dict):
    L, lang = LABELS[content["meta"]["language"]], content["meta"]["language"]
    _heading(doc, L["intro"], 1)
    for para in content["intro"]:
        _rich(doc.add_paragraph(), para)
    for i, ch in enumerate(content["chapters"]):
        _heading(doc, chapter_title(lang, i, ch["title"]), 1)
        for j, s in enumerate(ch["sections"]):
            _heading(doc, f"{i + 1}.{j + 1}. {s['title']}", 2)
            for para in s["paragraphs"]:
                _rich(doc.add_paragraph(), para)
    _heading(doc, L["conclusion"], 1)
    for para in content["conclusion"]:
        _rich(doc.add_paragraph(), para)
    _heading(doc, L["references"], 1)
    for n, ref in enumerate(content["references"], 1):
        p = doc.add_paragraph()
        p.add_run(f"{n}. {ref}")


def _numbers(section, fmt: dict):
    """Page numbers of the main section (the title page has its own footer with the city and year)."""
    where = fmt.get("page_numbers") or "bottom_center"
    for part in (section.header, section.footer):
        part.is_linked_to_previous = False
    part = section.header if where.startswith("top") else section.footer
    p = part.paragraphs[0]
    _plain(p, WD_ALIGN_PARAGRAPH.RIGHT if where.endswith("right") else WD_ALIGN_PARAGRAPH.CENTER)
    _field(p, "PAGE", "2")
    other = section.footer if where.startswith("top") else section.header
    other.paragraphs[0].text = ""


def write(content: dict, fmt: dict, path: Path, pages: dict[int, int] | None = None) -> Path:
    fmt = {**DEFAULT_FORMAT, **(fmt or {})}
    doc = Document()
    lang = LABELS[content["meta"]["language"]]["lang"]
    _styles(doc, fmt, lang)
    _page_setup(doc.sections[0], fmt)
    _title_page(doc, content["meta"], fmt)
    main = doc.add_section(WD_SECTION.NEW_PAGE)
    _page_setup(main, fmt)
    _numbers(main, fmt)
    _toc(doc, content, fmt, pages)
    _body(doc, content)
    props = doc.core_properties
    props.title, props.author, props.subject = content["meta"]["topic"], content["meta"].get("student") or "", content["meta"].get("subject") or ""
    props.comments = props.keywords = ""
    doc.save(str(path))
    return path


# ---------------------------------------------------------------- page numbers from a real layout

def to_pdf(docx: Path, timeout: int = 120) -> Path:
    """LibreOffice, headless, with a profile of its own (a second soffice would otherwise wait for the first)."""
    profile = Path(tempfile.mkdtemp(prefix="slideai-lo-"))
    try:
        with _office:
            subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
                            "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                           check=True, capture_output=True, timeout=timeout)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    pdf = docx.with_suffix(".pdf")
    if not pdf.exists():
        raise RuntimeError("LibreOffice PDF yaratmadi")
    return pdf


def _key(text: str) -> str:
    return re.sub(r"[\W_]+", "", text).upper()


def pdf_pages(pdf: Path) -> list[str]:
    out = subprocess.run(["pdftotext", "-enc", "UTF-8", str(pdf), "-"], check=True, capture_output=True, timeout=60).stdout
    pages = out.decode("utf-8", "replace").split("\f")
    return pages[:-1] if pages and not pages[-1].strip() else pages


def find_pages(items: list[tuple[int, str]], pages: list[str]) -> dict[int, int]:
    """1-based page of each heading. The search starts after the title page and the MUNDARIJA (page 2) and only
    moves forward, so the MUNDARIJA's own lines are never taken for the headings. A part (level 1) starts a new page,
    so it must be the first thing on it ("Xulosa qilib aytganda" in the text is not the XULOSA)."""
    keys = [_key(p) for p in pages]
    found, at = {}, 2
    for k, (level, text) in enumerate(items):
        want = _key(text)[:40]
        for i in range(at, len(keys)):
            if want and (keys[i].lstrip("0123456789").startswith(want) if level == 1 else want in keys[i]):
                found[k], at = i + 1, i
                break
        else:
            found[k] = at + 1  # not found (odd characters): the page of the previous heading
    return found


def build(content: dict, fmt: dict, out_dir: Path, name: str = "referat") -> tuple[Path, Path | None, int | None]:
    """Write the file, lay it out, put the real page numbers into the MUNDARIJA. Without LibreOffice (or if it fails)
    the .docx is still made, its TOC marked to be refreshed by Word on opening; then there is no PDF."""
    docx = out_dir / f"{name}.docx"
    write(content, fmt, docx)
    try:
        first = pdf_pages(to_pdf(docx))
        numbers = find_pages(headings(content), first)
        write(content, fmt, docx, numbers)
        pdf = to_pdf(docx)
        return docx, pdf, len(pdf_pages(pdf))
    except (OSError, subprocess.SubprocessError, RuntimeError) as e:
        log.warning("referat layout failed, sending the .docx without page numbers: %s", e)
        _update_on_open(docx)
        return docx, None, None


def _update_on_open(docx: Path):
    doc = Document(str(docx))
    el = OxmlElement("w:updateFields"); el.set(qn("w:val"), "true")
    doc.settings.element.append(el)
    doc.save(str(docx))

