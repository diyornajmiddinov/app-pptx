"""HTTP routes not covered elsewhere: payment callbacks, the Telegram webhook, the desktop editor."""
import base64
import hashlib
import io

import httpx
import pytest
from conftest import add_user
from test_auth import init_data
from test_editor import png, sample_pptx

from app.api import main
from app.bot import handlers
from app.core.auth import make_editor_token
from app.core.config import settings
from app.services import billing, editor, payments

U = 701


@pytest.fixture
async def client():
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://t") as c:
        yield c


async def test_payme_endpoint_always_answers_json_rpc(client, monkeypatch):
    monkeypatch.setattr(settings, "payme_key", "k")
    r = await client.post("/api/pay/payme", content=b"{not json")
    assert r.status_code == 200 and r.json()["error"]["code"] == -32700
    await add_user(U)
    oid = await payments.create_order(U, 30000, "payme")
    auth = {"Authorization": "Basic " + base64.b64encode(b"Paycom:k").decode()}
    body = {"id": 1, "method": "CheckPerformTransaction", "params": {"amount": 3000000, "account": {"order_id": str(oid)}}}
    assert (await client.post("/api/pay/payme", json=body, headers=auth)).json()["result"] == {"allow": True}


async def test_click_endpoints_take_a_form_and_tell_the_user(client, tg, monkeypatch):
    monkeypatch.setattr(settings, "click_secret_key", "s")
    await add_user(U)
    oid = await payments.create_order(U, 50000, "click")

    def signed(f, prep=False):
        parts = [f["click_trans_id"], f["service_id"], "s", f["merchant_trans_id"], *([f["merchant_prepare_id"]] if prep else []),
                 f["amount"], f["action"], f["sign_time"]]
        return {**f, "sign_string": hashlib.md5("".join(parts).encode()).hexdigest()}
    base = {"click_trans_id": "9", "service_id": "1", "merchant_trans_id": str(oid), "amount": "50000", "sign_time": "t", "error": "0"}
    r = (await client.post("/api/pay/click/prepare", data=signed({**base, "action": "0"}))).json()
    assert r["error"] == 0
    r = (await client.post("/api/pay/click/complete", data=signed({**base, "action": "1", "merchant_prepare_id": str(r["merchant_prepare_id"])}, True))).json()
    assert r["error"] == 0 and await billing.balance(U) == 50000
    assert any("50" in t for t in tg.texts())  # "paid" message to the user


async def test_webhook_checks_the_secret(client, monkeypatch):
    seen = []

    async def handle(update):
        seen.append(update)
    monkeypatch.setattr(handlers, "handle", handle)
    monkeypatch.setattr(settings, "webhook_secret", "sekret")
    assert (await client.post("/telegram/webhook", json={"update_id": 1})).status_code == 403
    r = await client.post("/telegram/webhook", json={"update_id": 2}, headers={"X-Telegram-Bot-Api-Secret-Token": "sekret"})
    assert r.status_code == 200 and seen == [{"update_id": 2}]


async def test_health_and_templates(client):
    assert (await client.get("/health")).json()["templates"] == 40
    assert len((await client.get("/api/templates")).json()) == 40


# ---------------------------------------------------------------- the editor over HTTP

@pytest.fixture
def ed(tmp_path, monkeypatch):
    monkeypatch.setattr(editor, "DATA", tmp_path)
    return {"X-Editor-Token": make_editor_token(U)}


async def test_editor_round_trip(client, ed, tg):
    await add_user(U)
    assert (await client.get("/api/editor/docs")).status_code == 401
    link = (await client.get("/api/editor/link", headers={"X-Telegram-Init-Data": init_data(U)})).json()["url"]
    assert "/editor.html#t=" in link
    doc = (await client.post("/api/editor/import", content=sample_pptx(), headers={**ed, "X-Filename": "D%C9%99rs.pptx"})).json()["id"]
    assert (await client.get("/api/editor/docs", headers=ed)).json()["docs"][0]["name"] == "Dərs"
    m = (await client.get(f"/api/editor/docs/{doc}", headers=ed)).json()
    title = next(e for e in m["slides"][0]["elements"] if e["kind"] == "text")
    op = {"op": "text_runs", "slide": 1, "id": title["id"], "paragraphs": [{"runs": [{"t": "HTTP orqali"}]}]}
    r = await client.post(f"/api/editor/docs/{doc}/ops", headers=ed, json={"version": m["version"], "ops": [op], "full": False})
    assert r.json()["version"] == 2 and r.json()["can_undo"]
    assert (await client.post(f"/api/editor/docs/{doc}/ops", headers=ed, json={"version": 1, "ops": [op]})).status_code == 409
    assert (await client.post(f"/api/editor/docs/{doc}/undo", headers=ed)).json()["version"] == 1
    assert (await client.post(f"/api/editor/docs/{doc}/redo", headers=ed)).json()["version"] == 2
    up = (await client.post(f"/api/editor/docs/{doc}/upload", headers=ed, content=png())).json()["upload"]
    assert len(up) == 16
    pic = next(e for e in m["slides"][0]["elements"] if e["kind"] == "picture")
    media = await client.get(f"/api/editor/docs/{doc}/media/{pic['img']['src']}", params={"t": ed["X-Editor-Token"]})
    assert media.status_code == 200 and media.content[:4] == b"\x89PNG"
    dl = await client.get(f"/api/editor/docs/{doc}/download", headers=ed)
    from pptx import Presentation
    assert "HTTP orqali" in [s.text_frame.text for s in Presentation(io.BytesIO(dl.content)).slides[0].shapes if s.has_text_frame]
    assert (await client.post(f"/api/editor/docs/{doc}/rename", headers=ed, json={"name": "Yangi"})).json()["name"] == "Yangi"
    assert (await client.post(f"/api/editor/docs/{doc}/send", headers=ed)).status_code == 200
    assert tg.sent("sendDocument")
    assert (await client.post(f"/api/editor/docs/{doc}/delete", headers=ed)).json() == {"ok": True}
    assert (await client.get(f"/api/editor/docs/{doc}", headers=ed)).status_code == 404


async def test_editor_refuses_big_or_broken_uploads(client, ed, monkeypatch):
    monkeypatch.setattr(editor, "MAX_FILE", 1000)
    assert (await client.post("/api/editor/import", content=b"x" * 2000, headers=ed)).status_code == 413
    assert (await client.post("/api/editor/import", content=b"not a pptx", headers=ed)).status_code == 422
