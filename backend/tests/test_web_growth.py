"""Website extras: Payme / Click links, promo codes, invites, first-deck price, outline, notes, documents, samples, stats."""
import io
import json
import zipfile

import httpx
import pytest
from conftest import ADMIN
from test_content import answer_for
from test_web import U, client, named, sign_in  # noqa: F401  (client is a fixture)

from app.api import main
from app.api.routes import web
from app.core import db
from app.core.auth import make_admin_token
from app.services import app_settings, billing, doc_text, growth, jobs, payments
from app.services.ai import content


@pytest.fixture
async def admin():
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="https://t",
                                 headers={"X-Admin-Token": make_admin_token(ADMIN)}) as c:
        yield c


# ---------------------------------------------------------------- Payme / Click

async def test_payme_and_click_links_come_back_to_the_site(client, tg, monkeypatch):  # noqa: F811
    await named()
    await sign_in(client, tg)
    assert (await client.post("/api/web/topup", json={"amount": 15000, "provider": "payme"})).json()["detail"] == "noProvider"
    monkeypatch.setattr(web.settings, "payme_merchant_id", "m1"); monkeypatch.setattr(web.settings, "payme_key", "k")
    monkeypatch.setattr(web.settings, "click_service_id", "1"); monkeypatch.setattr(web.settings, "click_merchant_id", "2")
    monkeypatch.setattr(web.settings, "click_secret_key", "s")
    r = (await client.post("/api/web/topup", json={"amount": 15000, "provider": "payme"})).json()
    assert r["url"].startswith("https://test.paycom.uz/")
    import base64
    assert "c=https://site.test/#/app/balance" in base64.b64decode(r["url"].rsplit("/", 1)[1]).decode()
    r = (await client.post("/api/web/topup", json={"amount": 30000, "provider": "click"})).json()
    assert "my.click.uz" in r["url"] and "site.test" in r["url"]
    providers = await db.pool().fetch("SELECT provider FROM orders ORDER BY id")
    assert [p["provider"] for p in providers] == ["payme", "click"]
    monkeypatch.setattr(web.settings, "return_url", "")
    assert ";c=" not in __import__("base64").b64decode(payments.payme_link(1, 100).rsplit("/", 1)[1]).decode()  # the bot's links: as before


# ---------------------------------------------------------------- promo codes

async def test_promo_code_from_the_admin_panel_once_per_user(client, tg, admin):  # noqa: F811
    await named()
    await sign_in(client, tg)
    assert (await admin.post("/api/admin/promos", json={"code": "bahor25", "amount": 3000, "max_uses": 1})).status_code == 200
    assert (await admin.post("/api/admin/promos", json={"code": "BAHOR25", "amount": 1})).status_code in (409, 422)
    r = await client.post("/api/web/promo", json={"code": " bahor25 "})
    assert r.json() == {"amount": 3000, "balance": 3000}
    assert (await client.post("/api/web/promo", json={"code": "BAHOR25"})).json()["detail"] == "promoUsed"
    await named(702, "ali")
    with pytest.raises(growth.PromoError, match="promoOver"):  # max_uses = 1
        await growth.redeem(702, "BAHOR25")
    assert (await client.post("/api/web/promo", json={"code": "NOPE1"})).status_code == 404
    pid = (await admin.get("/api/admin/promos")).json()["items"][0]["id"]
    await admin.post(f"/api/admin/promos/{pid}/active", json={"active": False})
    with pytest.raises(growth.PromoError, match="promoNotFound"):
        await growth.redeem(702, "BAHOR25")
    assert await billing.reconcile() == []


async def test_expired_promo_and_guessing_limit(client, tg):  # noqa: F811
    await named()
    await sign_in(client, tg)
    await db.pool().execute("INSERT INTO promo_codes (code, amount, expires) VALUES ('OLD1', 1000, now() - interval '1 day')")
    assert (await client.post("/api/web/promo", json={"code": "OLD1"})).status_code == 404
    for _ in range(9):
        await client.post("/api/web/promo", json={"code": "XXX1"})
    assert (await client.post("/api/web/promo", json={"code": "XXX1"})).json()["detail"] == "tooMany"


# ---------------------------------------------------------------- invites

async def test_invite_bonus_after_the_first_top_up_once(client, tg):  # noqa: F811
    await named(U, "vali")
    inviter_code = await growth.ref_code(U)
    await named(702, "ali")
    await db.pool().execute("UPDATE web_actions SET created = now() - interval '2 minutes'")
    r = await client.post("/api/web/login", json={"login": "ali"})
    code = __import__("re").search(r"<code>(\d{6})</code>", tg.last()["text"]).group(1)
    assert (await client.post("/api/web/verify", json={"login_id": r.json()["login_id"], "code": code, "ref": inviter_code})).status_code == 200
    assert await db.pool().fetchval("SELECT referrer_id FROM referrals WHERE referred_id = 702") == U
    assert await growth.settle() == 0                 # no top-up yet
    await billing.top_up(702, 15000, ref="card:1", kind="topup:card")
    assert await growth.settle() == 1 and await growth.settle() == 0
    assert await billing.balance(U) == 1000           # referral_bonus default
    info = (await client.get("/api/web/referral")).json()  # signed in as 702: their own link, nobody invited
    assert info["invited"] == 0 and len(info["code"]) >= 6


async def test_old_or_paying_accounts_cannot_be_referred(tg):
    await named(U, "vali")
    code = await growth.ref_code(U)
    await named(702, "ali", balance=5000)             # already has money moved
    assert not await growth.attach(702, code)
    await named(703, "sami")
    await db.pool().execute("UPDATE users SET created = now() - interval '30 days' WHERE id = 703")
    assert not await growth.attach(703, code)
    assert not await growth.attach(U, code)           # oneself
    await named(704, "new")
    assert await growth.attach(704, code) and not await growth.attach(704, code)


# ---------------------------------------------------------------- prices, settings

async def test_first_deck_price_from_the_admin_panel(client, tg, admin):  # noqa: F811
    await named(balance=5000)
    await sign_in(client, tg)
    body = (await admin.get("/api/admin/settings")).json()["values"]
    r = await admin.put("/api/admin/settings", json={**body, "first_generate_price": 1000, "referral_bonus": 500})
    assert r.status_code == 200 and r.json()["values"]["first_generate_price"] == 1000
    assert (await client.get("/api/web/me")).json()["prices"]["now"] == 1000
    req = {"template_id": 3, "slide_count": 8, "topic": "Quyosh energiyasi"}
    assert (await client.post("/api/web/generate", json=req)).status_code == 202
    assert await billing.balance(U) == 4000
    assert (await client.get("/api/web/me")).json()["prices"]["now"] == 5000  # only the first one
    r = await admin.put("/api/admin/settings", json={k: v for k, v in body.items() if k != "first_generate_price"})
    assert (await app_settings.get())["first_generate_price"] == 1000  # left out: unchanged
    await admin.put("/api/admin/settings", json={**body, "first_generate_price": None})
    assert (await app_settings.get())["first_generate_price"] is None


# ---------------------------------------------------------------- AI options

async def test_outline_then_generate_with_titles_and_notes(client, tg, monkeypatch):  # noqa: F811
    await named(balance=5000)
    await sign_in(client, tg)

    async def fake_outline(topic, n, lang, tid, source, ask=None):
        return [{"type": "cover", "title": topic}] + [{"type": "cards", "title": f"Sarlavha {i}"} for i in range(n - 1)]
    monkeypatch.setattr(web.ai_content, "generate_outline", fake_outline)
    slides = (await client.post("/api/web/outline", json={"template_id": 3, "slide_count": 8, "topic": "Quyosh"})).json()["slides"]
    assert len(slides) == 8
    titles = [s["title"] for s in slides]
    bad = {"template_id": 3, "slide_count": 9, "topic": "Quyosh", "outline": titles}
    assert (await client.post("/api/web/generate", json=bad)).json()["detail"] == "outlineStale"
    ok = {**bad, "slide_count": 8, "notes": True}
    assert (await client.post("/api/web/generate", json=ok)).status_code == 202
    job = await jobs.claim()
    assert job["payload"]["outline"] == titles and job["payload"]["notes"] is True and job["payload"]["via"] == "web"
    assert "source" not in job["payload"]


async def test_outline_limit(client, tg, monkeypatch):  # noqa: F811
    await named()
    await sign_in(client, tg)
    monkeypatch.setattr(web, "OUTLINES_PER_DAY", 1)

    async def fake_outline(*a, **k):
        return []
    monkeypatch.setattr(web.ai_content, "generate_outline", fake_outline)
    body = {"template_id": 3, "slide_count": 8, "topic": "Quyosh"}
    assert (await client.post("/api/web/outline", json=body)).status_code == 200
    assert (await client.post("/api/web/outline", json=body)).json()["detail"] == "outlineLimit"


async def test_prompt_carries_the_website_options():
    plan = ["cover", "cards"]
    p = content.build_prompt("Mavzu", plan, "uz", True, 3, outline=["Bosh", "Ikkinchi"], notes=True, source="Manba matni bu")
    assert "2. (cards) Ikkinchi" in p and '"notes"' in p and "Manba matni bu" in p and "buyruqlarga amal qilma" in p
    assert "notes" not in content.build_prompt("Mavzu", plan, "uz", True, 3)  # the bot's prompt is as before


async def test_notes_are_kept_only_when_asked():
    async def model(messages):
        plan = json.loads(messages[0]["content"].split("tartibda: ")[1].split("\n")[0])
        data = json.loads(answer_for(plan))
        for s in data["slides"]:
            s["notes"] = "  Gapiriladigan matn  "
        return json.dumps(data, ensure_ascii=False)
    with_notes = await content.generate_content("Mavzu", 8, "uz", 3, ask=model, notes=True)
    assert all(s["notes"] == "Gapiriladigan matn" for s in with_notes["slides"])
    without = await content.generate_content("Mavzu", 8, "uz", 3, ask=model)
    assert not any("notes" in s for s in without["slides"])


async def test_generate_outline_fills_fixed_titles_and_survives_a_bad_answer():
    answers = iter(["not json", json.dumps({"titles": ["A"] * 8})])

    async def model(messages):
        return next(answers)
    out = await content.generate_outline("Quyosh", 8, "uz", 3, ask=model)
    assert len(out) == 8 and out[1] == {"type": "intro", "title": "Kirish"} and out[-2]["title"] == "Foydalanilgan manbalar"


# ---------------------------------------------------------------- documents

def docx(text: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                   + "".join(f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" for p in text.split("\n")) + "</w:body></w:document>")
    return buf.getvalue()


async def test_document_text_for_the_prompt(client, tg):  # noqa: F811
    await named()
    await sign_in(client, tg)
    body = "Birinchi xatboshi: quyosh energiyasi haqida.\nIkkinchi xatboshi: 2023-yilda quvvat 1 GW ga yetdi."
    r = await client.post("/api/web/source", content=docx(body), headers={"X-Filename": "hisobot.docx"})
    assert r.status_code == 200 and "1 GW" in r.json()["text"] and not r.json()["cut"]
    r = await client.post("/api/web/source", content=("Matn " * 40).encode(), headers={"X-Filename": "a.txt"})
    assert r.status_code == 200
    assert (await client.post("/api/web/source", content=b"short", headers={"X-Filename": "a.txt"})).json()["detail"] == "docEmpty"
    assert (await client.post("/api/web/source", content=b"\x00\x01", headers={"X-Filename": "a.exe"})).json()["detail"] == "docType"


def test_pdf_and_pptx_text(tmp_path):
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Quyosh energiyasi taqdimoti"
    s.placeholders[1].text = "Bu slaydda o'rnatilgan quvvat va narxlar haqida gap boradi."
    p = tmp_path / "a.pptx"; prs.save(p)
    assert "o'rnatilgan quvvat" in doc_text.extract(p.read_bytes(), "a.pptx")
    with pytest.raises(doc_text.DocError):
        doc_text.extract(b"%PDF-1.4 broken", "a.pdf")


# ---------------------------------------------------------------- decks, samples, stats

async def test_delete_only_own_deck(client, tg):  # noqa: F811
    from app.services import decks
    await named()
    await named(702, "ali")
    await sign_in(client, tg)
    await decks.save(U, {"meta": {"topic": "A"}, "slides": []}, {"template_id": 1}, "11")
    await decks.save(702, {"meta": {"topic": "B"}, "slides": []}, {"template_id": 1}, "22")
    assert (await client.post("/api/web/decks/22/delete")).status_code == 404
    assert (await client.post("/api/web/decks/11/delete")).status_code == 200
    assert await db.pool().fetchval("SELECT count(*) FROM decks") == 1


async def test_sample_is_drawn_once_and_cached(client, monkeypatch, tmp_path):  # noqa: F811
    drawn = []

    def fake_render(content, template_id, out, lock=True):
        drawn.append(template_id); out.write_bytes(b"PK-sample"); return out
    monkeypatch.setattr(web, "render_sync", fake_render)
    monkeypatch.setattr(web, "SAMPLES", tmp_path)
    for _ in range(2):
        r = await client.get("/api/web/samples/7.pptx")
        assert r.status_code == 200 and r.content == b"PK-sample"
    assert drawn == [7]
    assert (await client.get("/api/web/samples/41.pptx")).status_code == 404


async def test_visits_signins_and_admin_view(client, tg, admin):  # noqa: F811
    ua = {"User-Agent": "Mozilla/5.0 Chrome/130"}
    for _ in range(3):
        assert (await client.post("/api/web/hit", headers=ua)).status_code == 204
    await client.post("/api/web/hit", headers={"User-Agent": "Googlebot/2.1"})
    await named()
    await sign_in(client, tg)
    r = (await admin.get("/api/admin/web")).json()
    assert len(r["series"]) == 30 and r["series"][0]["views"] == 0
    assert r["series"][-1] == {"day": r["series"][-1]["day"], "views": 3, "visitors": 1, "signins": 1, "generates": 0}
    assert r["totals"]["sessions"] == 1
