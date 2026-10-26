"""The desktop editor: import, the browser model (pptx_model), every edit operation, versions, undo / redo, safety."""
import io
import shutil

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu, Inches

from app.services import editor

U = 900


def png(color="red", size=(200, 100)) -> bytes:
    b = io.BytesIO(); Image.new("RGB", size, color).save(b, "PNG"); return b.getvalue()


def sample_pptx() -> bytes:
    """Two slides: a title, a text box, a coloured rectangle, a line, a table, a picture; notes on slide 1."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])          # title only
    s.shapes.title.text = "Asosiy sarlavha"
    tb = s.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    tb.text_frame.text = "Birinchi qator"
    tb.text_frame.add_paragraph().text = "Ikkinchi qator"
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6), Inches(2), Inches(2), Inches(1))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x2A, 0x78, 0xD6)
    s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(4), Inches(8), Inches(4))
    tbl = s.shapes.add_table(2, 2, Inches(1), Inches(4.5), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(1, 1).text = "A1", "B2"
    s.shapes.add_picture(io.BytesIO(png()), Inches(6), Inches(4.5), Inches(2), Inches(1))
    s.notes_slide.notes_text_frame.text = "Eslatma"
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "Ikkinchi slayd"
    out = io.BytesIO(); prs.save(out); return out.getvalue()


@pytest.fixture(autouse=True)
def data_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(editor, "DATA", tmp_path)  # never the real data/


@pytest.fixture
def doc():
    return editor.import_file(U, sample_pptx(), "Mening taqdimotim.pptx")


def by_kind(model, slide=0):
    out: dict[str, list] = {}
    for e in model["slides"][slide]["elements"]:
        out.setdefault(e["kind"], []).append(e)
    return out


def hexof(color: str) -> str:
    return color.upper().lstrip("#")[:6]


def text_of(e):
    return "\n".join("".join(r.get("t", "") for r in p["runs"]) for p in e["text"]["paragraphs"])


def edit(doc, *ops):
    d = editor.describe(U, doc)
    return editor.apply(U, doc, list(ops), d["version"])


# ---------------------------------------------------------------- import and the model

def test_import_rejects_what_is_not_a_presentation():
    for data, name, err in ((b"x", "a.pdf", "badFile"), (b"not a zip", "a.pptx", "badFile"), (b"0" * (editor.MAX_FILE + 1), "a.pptx", "tooBig")):
        with pytest.raises(editor.EditorError, match=err):
            editor.import_file(U, data, name)


def test_model_has_every_drawable_element(doc):
    m = editor.describe(U, doc)
    assert m["name"] == "Mening taqdimotim" and len(m["slides"]) == 2 and m["version"] == 1 and not m["can_undo"]
    k = by_kind(m)
    texts = sorted(text_of(e) for e in k["text"])
    assert texts == ["Asosiy sarlavha", "Birinchi qator\nIkkinchi qator"]
    assert k["shape"][0]["fill"] and hexof(k["shape"][0]["fill"]["color"]) == "2A78D6"
    assert k["line"] and k["table"] and k["picture"]
    head = k["table"][0]["table"]["rows"][0]["cells"][0]
    run = head["paragraphs"][0]["runs"][0]
    assert run["t"] == "A1" and run["b"] and hexof(run["color"]) == "FFFFFF"   # the table style's header row ...
    assert hexof(head["fill"]["color"]) == "4F81BD"                           # ... resolved from the theme
    assert m["slides"][0]["notes"] == "Eslatma" and len(m["slides"][0]["layer"]) == 20
    assert editor.list_docs(U)[0]["id"] == doc


def test_paths_cannot_leave_the_users_folder(doc):
    for bad in ("../../etc", "0" * 15, doc.upper()):
        with pytest.raises(editor.EditorError, match="notFound"):
            editor.describe(U, bad)
    with pytest.raises(editor.EditorError):
        editor.describe(U + 1, doc)                      # another user's document
    with pytest.raises(editor.EditorError):
        editor.media_file(U, doc, "../meta.json")


# ---------------------------------------------------------------- operations

def test_text_move_style_and_the_saved_file(doc):
    title = by_kind(editor.describe(U, doc))["text"][0]
    m = edit(doc, {"op": "text_runs", "slide": 1, "id": title["id"],
                   "paragraphs": [{"src": 0, "align": "center", "runs": [{"t": "Yangi ", "b": True}, {"t": "sarlavha", "color": "#FF0000", "sz": 40}]}]},
             {"op": "move", "slide": 1, "id": title["id"], "x": 914400, "y": 457200, "w": 6000000, "h": 900000},
             {"op": "rotate", "slide": 1, "id": title["id"], "rot": 370})
    t = next(e for e in by_kind(m)["text"] if e["id"] == title["id"])
    assert text_of(t) == "Yangi sarlavha" and (t["x"], t["y"]) == (914400, 457200) and round(t["rot"]) == 10
    runs = t["text"]["paragraphs"][0]["runs"]
    assert runs[0]["b"] and hexof(runs[1]["color"]) == "FF0000" and runs[1]["sz"] == 40
    path, name = editor.current_file(U, doc)
    shape = next(s for s in Presentation(str(path)).slides[0].shapes if s.shape_id == title["id"])
    assert shape.text_frame.text == "Yangi sarlavha" and shape.left == Emu(914400) and name == "Mening taqdimotim"


def test_table_cells_fill_order_duplicate_delete(doc):
    k = by_kind(editor.describe(U, doc))
    table, box = k["table"][0], k["shape"][0]
    m = edit(doc, {"op": "cell_runs", "slide": 1, "id": table["id"], "r": 1, "c": 0, "paragraphs": [{"runs": [{"t": "Yangi katak"}]}]},
             {"op": "fill", "slide": 1, "id": box["id"], "color": "#00AA00"},
             {"op": "duplicate", "slide": 1, "id": box["id"]},
             {"op": "order", "slide": 1, "id": box["id"], "where": "front"})
    k = by_kind(m)
    assert k["table"][0]["table"]["rows"][1]["cells"][0]["paragraphs"][0]["runs"][0]["t"] == "Yangi katak"
    assert len(k["shape"]) == 2 and all(hexof(s["fill"]["color"]) == "00AA00" for s in k["shape"])
    assert m["slides"][0]["elements"][-1]["id"] == box["id"]          # brought to the front
    m = edit(doc, {"op": "delete", "slide": 1, "id": box["id"]})
    assert len(by_kind(m)["shape"]) == 1


def test_slides_add_duplicate_move_hide_delete_bg_notes(doc):
    m = edit(doc, {"op": "slide_duplicate", "slide": 1}, {"op": "slide_add", "after": 3}, {"op": "slide_move", "slide": 4, "to": 1},
             {"op": "slide_hide", "slide": 2, "hidden": True}, {"op": "bg", "slide": 3, "color": "#112233"},
             {"op": "notes", "slide": 3, "text": "Yangi eslatma"})
    assert len(m["slides"]) == 4 and m["slides"][1]["hidden"] and hexof(m["slides"][2]["bg"]) == "112233"
    assert m["slides"][2]["notes"] == "Yangi eslatma"
    assert sorted(text_of(e) for e in by_kind(m, 2)["text"]) == ["Asosiy sarlavha", "Birinchi qator\nIkkinchi qator"]  # the copy
    for _ in range(3):
        m = edit(doc, {"op": "slide_delete", "slide": 1})
    with pytest.raises(editor.EditorError, match="lastSlide"):
        edit(doc, {"op": "slide_delete", "slide": 1})


def test_images_are_checked_added_and_replaced(doc):
    with pytest.raises(editor.EditorError, match="badImage"):
        editor.save_upload(U, doc, b"not an image")
    up = editor.save_upload(U, doc, png("blue", (400, 300)))
    pic = by_kind(editor.describe(U, doc))["picture"][0]
    m = edit(doc, {"op": "add_image", "slide": 2, "upload": up}, {"op": "image", "slide": 1, "id": pic["id"], "upload": up})
    assert by_kind(m, 1)["picture"]
    assert by_kind(m)["picture"][0]["img"]["src"] != pic["img"]["src"]
    assert editor.media_file(U, doc, by_kind(m)["picture"][0]["img"]["src"]).exists()
    with pytest.raises(editor.EditorError, match="badImage"):
        edit(doc, {"op": "add_image", "slide": 1, "upload": "0" * 16})


def test_bad_operations_change_nothing(doc):
    v = editor.describe(U, doc)["version"]
    for op in ({"op": "explode", "slide": 1, "id": 2}, {"op": "move", "slide": 9, "id": 2}, {"op": "cell", "slide": 1, "id": 2, "r": 0, "c": 0}):
        with pytest.raises(editor.EditorError):
            editor.apply(U, doc, [op], v)
    assert editor.describe(U, doc)["version"] == v


# ---------------------------------------------------------------- versions

def test_undo_redo_and_a_new_change_drops_the_redo_branch(doc):
    title = by_kind(editor.describe(U, doc))["text"][0]

    def say(text):
        return edit(doc, {"op": "text_runs", "slide": 1, "id": title["id"], "paragraphs": [{"runs": [{"t": text}]}]})
    say("ikki"); say("uch")
    m = editor.step(U, doc, -1)
    assert text_of(next(e for e in by_kind(m)["text"] if e["id"] == title["id"])) == "ikki" and m["can_redo"]
    m = editor.step(U, doc, +1)
    assert text_of(next(e for e in by_kind(m)["text"] if e["id"] == title["id"])) == "uch"
    editor.step(U, doc, -1)
    m = say("to'rt")
    assert not m["can_redo"]


def test_stale_version_is_refused_and_old_versions_are_trimmed(doc, monkeypatch):
    monkeypatch.setattr(editor, "MAX_VERSIONS", 3)
    title = by_kind(editor.describe(U, doc))["text"][0]
    for i in range(5):
        edit(doc, {"op": "text_runs", "slide": 1, "id": title["id"], "paragraphs": [{"runs": [{"t": str(i)}]}]})
    folder = editor._folder(U, doc)
    assert len(list(folder.glob("v*.pptx"))) == 3
    with pytest.raises(editor.EditorError, match="stale"):
        editor.apply(U, doc, [{"op": "rotate", "slide": 1, "id": title["id"], "rot": 5}], 1)


def test_rename_and_delete(doc):
    assert editor.rename(U, doc, 'Yangi: nom/1')["name"] == "Yangi_ nom_1"
    editor.delete_doc(U, doc)
    assert editor.list_docs(U) == []


@pytest.mark.slow
def test_background_layer_is_drawn_without_the_elements(doc):
    if not shutil.which("soffice"):
        pytest.skip("LibreOffice not installed")
    m = editor.describe(U, doc)
    f = editor.layer_file(U, doc, m["slides"][0]["layer"])
    with Image.open(f) as im:
        assert im.width == 1600
