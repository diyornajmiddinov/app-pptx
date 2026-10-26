"""Every template draws the sample deck into a .pptx that opens, with every slide (node + the engine)."""
import json
import shutil

import pytest
from pptx import Presentation

from app.core.config import ENGINE_DIR
from app.services.ai.render import render, render_sync

pytestmark = pytest.mark.slow
SAMPLE = json.loads((ENGINE_DIR / "content" / "soliq.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def sample():
    if not shutil.which("node") or not (ENGINE_DIR / "node_modules").exists():
        pytest.skip("node / engine dependencies not installed")
    c = json.loads(json.dumps(SAMPLE))
    c["meta"] = {"topic": "Soliq islohotlari", "author": "Test", "org": "TDIU", "loc": "Toshkent", "year": "2026"}
    for s in c["slides"]:
        if isinstance(s.get("image"), str) and not s["image"].startswith("/"):
            s["image"] = str(ENGINE_DIR / s["image"])
    return c


@pytest.mark.parametrize("template_id", range(1, 41))
def test_template_renders(template_id, sample, tmp_path):
    out = render_sync(sample, template_id, tmp_path / "t.pptx")
    prs = Presentation(out)
    assert len(prs.slides) == len(sample["slides"])


async def test_async_render_and_a_broken_deck(sample, tmp_path):
    import subprocess
    out = await render(sample, 7, tmp_path / "a.pptx")
    assert len(Presentation(out).slides) == len(sample["slides"])
    with pytest.raises(subprocess.CalledProcessError) as e:
        await render({"slides": "broken", "meta": {}}, 3, tmp_path / "b.pptx")
    assert "TypeError" in e.value.stderr  # what alerts.why shows to the admins
